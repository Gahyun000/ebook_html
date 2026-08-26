import base64
import json
import os
import re
import subprocess
import sys
import time
import pathlib
import urllib.request
import urllib.error
from typing import List, Optional

from fastapi import FastAPI, Response, UploadFile, File, Form
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import StreamingResponse
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel

from server.settings_store import (
    init_db, load_llm_settings, save_llm_settings, mask_key, llm_endpoint, llm_configured,
)
from server import chat as chat_engine
from server import projects as projects_store
from server import notes as notes_store

HERE = pathlib.Path(__file__).resolve().parent
EBOOK_HTML = HERE.parent
GIT_ROOT = EBOOK_HTML.parent
UNIEVER = pathlib.Path(os.environ.get("UNIEVER_EBOOK") or (GIT_ROOT / "uniever_ebook"))
GEN_DIR = UNIEVER / "ebook-generator"
GEN_PY = GEN_DIR / "generator.py"
EBOOKS = UNIEVER / "ebooks"

app = FastAPI(title="ebook_html M2")
app.add_middleware(CORSMiddleware, allow_origins=["*"], allow_methods=["*"], allow_headers=["*"])


class Page(BaseModel):
    role: str
    seq: Optional[int] = None
    title: str = ""
    imageDataUrl: str


class Hotspot(BaseModel):
    seq: int
    x: float
    y: float
    w: float
    h: float
    label: str = ""


class BuildReq(BaseModel):
    title: str = "제목 없음"
    orientation: str = "portrait"
    theme: str = "light"
    pages: List[Page]
    hotspots: List[Hotspot] = []
    project_id: Optional[str] = None


def sanitize(name: str) -> str:
    s = re.sub(r'[\\/:*?"<>|\n\r\t]+', " ", name or "").strip()
    return s or "page"


def dataurl_bytes(d: str) -> bytes:
    b64 = d.split(",", 1)[1] if "," in d else d
    return base64.b64decode(b64)


@app.get("/api/health")
def health():
    return {"ok": True, "generator": str(GEN_PY), "generator_exists": GEN_PY.exists(), "ebooks_dir": str(EBOOKS)}


@app.post("/api/build")
def build(req: BuildReq):
    if not GEN_PY.exists():
        return {"ok": False, "error": "generator.py 없음: %s" % GEN_PY}
    ts = time.strftime("%Y%m%d_%H%M%S")
    work = EBOOK_HTML / "_build_input" / ts
    work.mkdir(parents=True, exist_ok=True)
    n = 0
    for p in req.pages:
        try:
            data = dataurl_bytes(p.imageDataUrl)
        except Exception as e:
            return {"ok": False, "error": "이미지 디코드 실패: %s" % e}
        if p.role == "cover":
            fn = "표지.png"
        elif p.role == "back":
            fn = "뒷표지.png"
        elif p.role == "toc":
            fn = "00. 목차.png"
        else:
            n += 1
            seq = p.seq if p.seq is not None else n
            fn = "%02d. %s.png" % (seq, sanitize(p.title))
        (work / fn).write_bytes(data)
    # 목차 클릭 영역: ebook_html이 정확한 좌표를 계산해 보냈으면 그대로 사이드카로 저장한다.
    # generator.py 는 _hotspots.json 을 OCR 자동추정보다 우선한다(정확·확정적).
    if req.hotspots:
        payload = {"hotspots": [h.model_dump() for h in req.hotspots]}
        (work / "_hotspots.json").write_text(
            json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8"
        )
    # 통짜(full-bleed) 렌더: ebook_html이 이미 완성된 EVER-PEAK 페이지를 통째로 굽기 때문에
    # generator가 위에 제목 헤더(card 스타일)를 얹지 않도록 image + full-bleed 로 만든다.
    proc = subprocess.run(
        [sys.executable, str(GEN_PY), "build", str(work), "--style", "image", "--full-bleed", "--title", req.title],
        cwd=str(GEN_DIR), capture_output=True, text=True,
    )
    if proc.returncode != 0:
        return {"ok": False, "error": (proc.stderr or proc.stdout)[-2000:]}
    m = re.search(r"생성됨:\s*(.+)", proc.stdout)
    out_path = m.group(1).strip() if m else None
    eid = os.path.basename(out_path) if out_path else None
    if req.project_id and eid:
        try:
            projects_store.set_published(req.project_id, eid)
        except Exception:
            pass
    return {"ok": True, "id": eid, "path": out_path, "url": ("/ebooks/%s/index.html" % eid) if eid else None, "log": proc.stdout[-1500:]}


# ───────────────────────── HTML → 덱 변환 (html_pdf_agent 이식) ─────────────────────────
@app.post("/api/deck")
async def deck(file: UploadFile = File(...), theme: str = Form("light"), summarize: bool = Form(False)):
    """HTML 가져오기 → 편집 가능한 덱 PPTX(+PDF·썸네일). 결과 파일은 /deck_out 로 서빙."""
    try:
        from server.deck.convert import convert
    except Exception as e:  # noqa: BLE001
        return {"ok": False, "error": "덱 변환 모듈 로드 실패: %s" % e}
    raw = await file.read()
    html = raw.decode("utf-8", "ignore")
    # LLM 이 연결돼 있으면 섹션을 헤드라인+핵심 몇 줄로 요약해 성근 슬라이드로(없으면 휴리스틱 압축 폴백).
    summarize_fn = None
    if summarize and load_llm_settings().get("configured"):
        def summarize_fn(title: str, text: str):
            try:
                content = _llm_chat([
                    {"role": "system", "content": _SUMM_SYS},
                    {"role": "user", "content": "[제목] %s\n[내용]\n%s\n\n위 섹션을 슬라이드 1장으로 요약해 JSON으로만 답하라." % (
                        title, (text or "")[:4000])},
                ], max_tokens=400, temperature=0.2)
                obj = _extract_json(content) or {}
                bl = [str(b).strip() for b in (obj.get("bullets") or []) if str(b).strip()][:4]
                return {"headline": str(obj.get("headline") or "")[:40], "bullets": bl} if bl else None
            except Exception:
                return None
    try:
        res = convert(html, theme, file.filename or "deck.html", summarize_fn=summarize_fn)
    except Exception as e:  # noqa: BLE001
        return {"ok": False, "error": str(e)}
    tok = res["token"]
    return {
        "ok": True,
        "pages": res["pages"],
        "pptx_url": "/deck_out/%s/%s" % (tok, res["pptx"]),
        "pdf_url": ("/deck_out/%s/%s" % (tok, res["pdf"])) if res["pdf"] else None,
        "thumbs": ["/deck_out/%s/%s" % (tok, t) for t in res["thumbs"]],
        "ir": res.get("ir"),   # 편집 요소 재현(구글 슬라이드식)용 덱 IR
    }


def _deck_summarizer():
    """LLM 설정이 있으면 섹션 요약기 반환(없으면 None). /api/deck/stream 공용."""
    if not load_llm_settings().get("configured"):
        return None

    def summarize_fn(title: str, text: str):
        try:
            content = _llm_chat([
                {"role": "system", "content": _SUMM_SYS},
                {"role": "user", "content": "[제목] %s\n[내용]\n%s\n\n위 섹션을 슬라이드 1장으로 요약해 JSON으로만 답하라." % (
                    title, (text or "")[:4000])},
            ], max_tokens=400, temperature=0.2)
            obj = _extract_json(content) or {}
            bl = [str(b).strip() for b in (obj.get("bullets") or []) if str(b).strip()][:4]
            return {"headline": str(obj.get("headline") or "")[:40], "bullets": bl} if bl else None
        except Exception:
            return None
    return summarize_fn


@app.post("/api/deck/stream")
async def deck_stream(file: UploadFile = File(...), theme: str = Form("light"), summarize: bool = Form(False)):
    """HTML → 덱 변환을 SSE 로. 추출/빌드/PDF/슬라이드별 렌더 진행을 실시간 전송(data: JSON)."""
    raw = await file.read()
    html = raw.decode("utf-8", "ignore")
    name = file.filename or "deck.html"
    summarize_fn = _deck_summarizer() if summarize else None

    def gen():
        try:
            from server.deck.convert import convert_stream
            for ev in convert_stream(html, theme, name, summarize_fn=summarize_fn):
                yield ("data: " + json.dumps(ev, ensure_ascii=False) + "\n\n").encode("utf-8")
        except Exception as e:  # noqa: BLE001
            yield ("data: " + json.dumps({"stage": "error", "error": str(e)}) + "\n\n").encode("utf-8")

    return StreamingResponse(gen(), media_type="text/event-stream",
                             headers={"Cache-Control": "no-cache", "X-Accel-Buffering": "no"})


# ───────────────────────── 환경설정: LLM (Agentic-PM 이식) ─────────────────────────
init_db()


class LlmSettingsIn(BaseModel):
    provider: Optional[str] = None
    base_url: Optional[str] = None
    user_id: Optional[str] = None
    api_key: Optional[str] = None
    model: Optional[str] = None
    enabled: Optional[bool] = None
    timeout: Optional[float] = None


def _effective_payload(s: dict) -> dict:
    return {
        "effective": {
            "provider": s.get("provider"),
            "base_url": s.get("base_url"),
            "user_id": s.get("user_id"),
            "model": s.get("model"),
            "enabled": s.get("enabled"),
            "timeout": s.get("timeout"),
            "api_key_masked": mask_key(s.get("api_key") or ""),
            "configured": s.get("configured"),
        },
        "providers": ["self", "openai", "anthropic", "custom"],
    }


def _coalesce(incoming: dict, current: dict, key: str, default=None):
    if key in incoming and incoming[key] not in (None, ""):
        return incoming[key]
    val = current.get(key)
    return val if val not in (None, "") else default


@app.get("/api/settings/llm")
def get_llm():
    return _effective_payload(load_llm_settings())


@app.put("/api/settings/llm")
def put_llm(payload: LlmSettingsIn):
    save_llm_settings(payload.model_dump(exclude_unset=True), updated_by="local")
    return _effective_payload(load_llm_settings())


def _llm_test_call(provider: str, url: str, user_id: str, api_key: str, model: str, timeout: float) -> dict:
    """설정값으로 LLM을 실제 1회 호출해 연결을 확인(OpenAI 호환 / Anthropic)."""
    headers = {"Content-Type": "application/json"}
    if provider == "anthropic":
        headers["x-api-key"] = api_key
        headers["anthropic-version"] = "2023-06-01"
        body = {"model": model, "max_tokens": 32, "messages": [{"role": "user", "content": "연결 테스트라고 한 문장으로 답하세요."}]}
    else:
        headers["Authorization"] = f"Bearer {api_key}"
        headers["x-api-key"] = api_key
        if user_id:
            headers["X-User-Id"] = user_id
        body = {"model": model, "max_tokens": 32,
                "messages": [{"role": "system", "content": "연결 테스트입니다. 짧게 답하세요."},
                             {"role": "user", "content": "연결 테스트라고 한 문장으로 답하세요."}]}
        # 유니에버 내부 게이트웨이 호환: user_id·api_key 를 본문에도 넣는다(헤더만 요구하지 않는 경우 대비).
        if user_id:
            body["user_id"] = user_id
        if api_key:
            body["api_key"] = api_key
    data = json.dumps(body).encode("utf-8")
    req = urllib.request.Request(url, data=data, headers=headers, method="POST")
    started = time.perf_counter()
    with urllib.request.urlopen(req, timeout=timeout) as resp:
        raw = resp.read().decode("utf-8", "replace")
        latency = int((time.perf_counter() - started) * 1000)
        sample = ""
        try:
            j = json.loads(raw)
            if provider == "anthropic":
                sample = (j.get("content", [{}])[0] or {}).get("text", "")
            else:
                sample = (((j.get("choices") or [{}])[0]).get("message") or {}).get("content", "")
        except Exception:
            sample = raw[:200]
        return {"status_code": resp.status, "latency_ms": latency, "sample": (sample or "")[:200]}


@app.post("/api/settings/llm/test")
def test_llm(payload: LlmSettingsIn):
    current = load_llm_settings()
    inc = payload.model_dump(exclude_unset=True)
    provider = str(_coalesce(inc, current, "provider", "self") or "self").lower()
    base_url = str(_coalesce(inc, current, "base_url", "") or "")
    user_id = str(_coalesce(inc, current, "user_id", "") or "")
    api_key = str(_coalesce(inc, current, "api_key", "") or "")
    model = str(_coalesce(inc, current, "model", "") or "")
    enabled = bool(_coalesce(inc, current, "enabled", True))
    timeout = float(_coalesce(inc, current, "timeout", 45) or 45)
    url = llm_endpoint(provider, base_url)
    configured = bool(url) and llm_configured(provider, base_url, api_key, model, user_id, enabled)
    result = {
        "ok": False, "configured": configured, "provider": provider, "url": url, "model": model,
        "user_id_set": bool(user_id), "api_key_set": bool(api_key),
        "status_code": None, "latency_ms": None, "message": "", "sample": "",
    }
    if not configured:
        result["message"] = "Base URL, User ID, API Key, 모델명이 모두 필요하고 LLM이 활성화되어야 합니다."
        return result
    try:
        r = _llm_test_call(provider, url, user_id, api_key, model, min(timeout, 30.0))
        result.update(r)
        result["ok"] = (r["status_code"] or 500) < 500
        result["message"] = "연결 성공" if result["ok"] else f"서버 오류: HTTP {r['status_code']}"
    except urllib.error.HTTPError as e:
        result["status_code"] = e.code
        result["message"] = f"연결 실패: HTTP {e.code}"
    except Exception as e:  # noqa: BLE001
        result["message"] = f"연결 실패: {e}"
    return result


# ───────────────────────── 채팅 (Agentic-PM chat_engine 코어 이식·중립화) ─────────────────────────
class ScreenCtx(BaseModel):
    page: Optional[str] = None
    selected_project_id: Optional[str] = None


class ChatIn(BaseModel):
    message: str = ""
    session_id: Optional[str] = None
    confirm: Optional[bool] = None
    confirm_action_id: Optional[str] = None
    screen_context: Optional[ScreenCtx] = None
    # G1 문서 컨텍스트 계약 — 프론트가 매 턴 보내는 현재 이북 스냅샷.
    # {title, orientation, theme, selectedPageId, pages:[{id, cardKey, title, summary}]}
    # 형태 드리프트에 관대하도록 permissive dict 로 받는다(소비는 G3 생성 레인에서).
    book_state: Optional[dict] = None


@app.post("/api/chat/v2")
def chat_v2(req: ChatIn):
    return chat_engine.respond(req.message, req.session_id, req.confirm, req.confirm_action_id, req.book_state)


class ResetIn(BaseModel):
    session_id: Optional[str] = None


@app.post("/api/chat/reset")
def chat_reset(req: ResetIn):
    return chat_engine.reset_session(req.session_id)


@app.get("/api/chat/conversations")
def chat_conversations():
    return {"conversations": chat_engine.list_conversations()}


@app.get("/api/chat/conversations/{cid}")
def chat_conversation(cid: str):
    return chat_engine.get_conversation(cid)


@app.delete("/api/chat/conversations/{cid}")
def chat_conversation_delete(cid: str):
    return chat_engine.delete_conversation(cid)


@app.post("/api/chat/v2/stream")
def chat_v2_stream(req: ChatIn):
    return StreamingResponse(
        chat_engine.sse_stream(req.message, req.session_id, req.confirm, req.confirm_action_id, req.book_state),
        media_type="text/event-stream",
        headers={"Cache-Control": "no-cache", "X-Accel-Buffering": "no"},
    )


class DocxIn(BaseModel):
    title: str = "답변"
    content: str = ""


@app.post("/api/export/docx")
def export_docx(req: DocxIn):
    try:
        from docx import Document  # python-docx
    except Exception:
        return {"ok": False, "error": "python-docx 미설치"}
    import io
    doc = Document()
    doc.add_heading(req.title or "답변", level=1)
    for line in (req.content or "").split("\n"):
        t = line.rstrip()
        if t.startswith("### "):
            doc.add_heading(t[4:], level=3)
        elif t.startswith("## "):
            doc.add_heading(t[3:], level=2)
        elif t.startswith("# "):
            doc.add_heading(t[2:], level=1)
        elif t.strip():
            doc.add_paragraph(t)
    buf = io.BytesIO()
    doc.save(buf)
    return Response(
        content=buf.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": 'attachment; filename="ebook_html-answer.docx"'},
    )


# ───────────────────────── (B) 섹션 → LLM 요약 → EVER-PEAK 카드 ─────────────────────────
class SummSection(BaseModel):
    title: str = ""
    text: str = ""


class SummReq(BaseModel):
    sections: List[SummSection]


def _llm_chat(messages: list, max_tokens: int = 512, temperature: float = 0.2) -> str:
    """저장된 LLM 설정으로 게이트웨이를 1회 호출해 텍스트를 받는다(요약용).
    api_key 는 서버 설정(DB/env)에서만 읽는다 — 코드/응답에 노출하지 않는다."""
    s = load_llm_settings()
    provider = str(s.get("provider") or "self").lower()
    base_url = str(s.get("base_url") or "")
    user_id = str(s.get("user_id") or "")
    api_key = str(s.get("api_key") or "")
    model = str(s.get("model") or "")
    url = llm_endpoint(provider, base_url)
    if not (url and model):
        raise RuntimeError("LLM 미설정 — 환경설정에서 base_url·model 을 넣으세요.")
    headers = {"Content-Type": "application/json"}
    body: dict = {"model": model, "messages": messages, "max_tokens": max_tokens, "temperature": temperature}
    if provider == "anthropic":
        headers["x-api-key"] = api_key
        headers["anthropic-version"] = "2023-06-01"
    else:
        if api_key:
            headers["Authorization"] = f"Bearer {api_key}"
            headers["x-api-key"] = api_key
            body["api_key"] = api_key          # 유니에버 게이트웨이 호환(본문 포함)
        if user_id:
            headers["X-User-Id"] = user_id
            body["user_id"] = user_id
    data = json.dumps(body).encode("utf-8")
    req = urllib.request.Request(url, data=data, headers=headers, method="POST")
    timeout = float(s.get("timeout") or 45)
    with urllib.request.urlopen(req, timeout=min(timeout, 90.0)) as resp:
        raw = resp.read().decode("utf-8", "replace")
    j = json.loads(raw)
    if provider == "anthropic":
        return (j.get("content", [{}])[0] or {}).get("text", "") or ""
    return (((j.get("choices") or [{}])[0]).get("message") or {}).get("content", "") or ""


def _extract_json(text: str):
    t = (text or "").strip()
    if t.startswith("```"):
        t = re.sub(r"^```[a-zA-Z]*\s*", "", t).rstrip("`").strip()
    m = re.search(r"\{.*\}", t, re.S)
    if not m:
        return None
    try:
        return json.loads(m.group(0))
    except Exception:
        return None


_SUMM_SYS = (
    "너는 제조업 임원 보고용 슬라이드 요약가다. 주어진 섹션을 슬라이드 한 장 분량으로 압축한다. "
    "핵심만 남기고 군더더기·중복은 버린다. 표는 대표 항목 몇 개만 남긴다. "
    "반드시 JSON 하나만 출력한다(설명 금지). "
    '형식: {"headline": "한 줄 제목(20자 내외)", "bullets": ["핵심 문장 3~5개(각 40자 내외)"], "cardType": "content"}'
)


@app.post("/api/summarize")
def summarize(req: SummReq):
    results = []
    for sec in req.sections:
        user = "[제목] %s\n[내용]\n%s\n\n위 섹션을 슬라이드 1장으로 요약해 JSON으로만 답하라." % (
            sec.title, (sec.text or "")[:4000])
        try:
            content = _llm_chat([
                {"role": "system", "content": _SUMM_SYS},
                {"role": "user", "content": user},
            ])
            obj = _extract_json(content) or {}
            bullets = [str(b).strip() for b in (obj.get("bullets") or []) if str(b).strip()][:5]
            results.append({
                "title": sec.title,
                "headline": (str(obj.get("headline") or "").strip() or sec.title)[:80],
                "bullets": [b[:120] for b in bullets],
                "cardType": str(obj.get("cardType") or "content"),
            })
        except Exception as e:  # noqa: BLE001
            # 실패한 섹션은 bullets 비움 → 프론트가 원문(A) 카드 그대로 유지
            results.append({"title": sec.title, "headline": sec.title, "bullets": [], "cardType": "content", "error": str(e)[:200]})
    return {"ok": True, "results": results}


# ───────────────────────── (G3) 브리프 → BookPlan 생성 하네스 ─────────────────────────
class PlanIn(BaseModel):
    brief: str = ""
    book_state: Optional[dict] = None
    retries: int = 2


@app.post("/api/plan")
def plan(req: PlanIn):
    """브리프 → 검증된 BookPlan(JSON). 서버 설정 LLM을 llm_fn 으로 감싸 planner 하네스에 주입한다.
    적용(카드 생성)은 G4 프론트 브리지에서. 여기서는 계획만 만든다."""
    from server.intent import planner

    def llm_fn(msgs: list) -> Optional[str]:
        try:
            return _llm_chat(msgs, max_tokens=1500, temperature=0.3)
        except Exception:
            return None

    return planner.make_plan(req.brief, llm_fn, req.book_state, retries=req.retries)


# ───────────────────────── (G5) 편집 명령 → 부분 수정 계획(edits/adds) ─────────────────────────
class EditIn(BaseModel):
    message: str = ""
    book_state: Optional[dict] = None
    retries: int = 2


@app.post("/api/edit")
def edit(req: EditIn):
    """편집 명령(이 장 다듬어/톤 통일/N장 추가) → 검증된 편집 계획. detect_edit 로 op 판정 후 make_edits.
    적용(스토어 반영)은 프론트 브리지(apply_page_edits)에서. 여기서는 계획만 만든다."""
    from server.intent import editor

    op = editor.detect_edit(req.message, req.book_state)
    if not op:
        return {"ok": False, "reason": "edit_not_detected", "op": None,
                "edits": [], "adds": [], "warnings": [], "summary": "편집 명령으로 인식하지 못했어요."}

    def llm_fn(msgs: list) -> Optional[str]:
        try:
            return _llm_chat(msgs, max_tokens=1500, temperature=0.3)
        except Exception:
            return None

    return editor.make_edits(op, req.message, llm_fn, req.book_state, retries=req.retries)



# ───────────────────────── 프로젝트(내 이북) ─────────────────────────
class ProjectCreateIn(BaseModel):
    name: Optional[str] = None
    state: Optional[dict] = None


class ProjectSaveIn(BaseModel):
    state: dict
    name: Optional[str] = None


class ProjectRenameIn(BaseModel):
    name: str


class VersionSaveIn(BaseModel):
    state: dict
    label: Optional[str] = None
    pinned: bool = False
    auto: bool = True


class VersionPatchIn(BaseModel):
    label: Optional[str] = None
    pinned: Optional[bool] = None


@app.get("/api/projects")
def projects_list():
    return {"projects": projects_store.list_projects()}


@app.post("/api/projects")
def projects_create(req: ProjectCreateIn):
    return projects_store.create_project(req.name, req.state)


@app.get("/api/projects/{pid}")
def projects_get(pid: str):
    p = projects_store.get_project(pid)
    if not p:
        return Response(status_code=404, content="not found")
    return p


@app.put("/api/projects/{pid}")
def projects_save(pid: str, req: ProjectSaveIn):
    return projects_store.save_project(pid, req.state, req.name)


@app.patch("/api/projects/{pid}")
def projects_rename(pid: str, req: ProjectRenameIn):
    return projects_store.rename_project(pid, req.name)


@app.delete("/api/projects/{pid}")
def projects_delete(pid: str):
    return projects_store.delete_project(pid)


@app.post("/api/projects/{pid}/duplicate")
def projects_duplicate(pid: str):
    d = projects_store.duplicate_project(pid)
    if not d:
        return Response(status_code=404, content="not found")
    return d


class NoteIn(BaseModel):
    project_id: str
    id: str
    title: str = ""
    blocks: list = []
    pinned: bool = False
    sort: float = 0


@app.get("/api/notes")
def notes_list(project_id: str):
    return {"notes": notes_store.list_notes(project_id)}


@app.post("/api/notes")
def notes_upsert(req: NoteIn):
    return notes_store.upsert_note(req.project_id, req.id, req.title, req.blocks, req.pinned, req.sort)


@app.delete("/api/notes/{nid}")
def notes_delete(nid: str):
    return notes_store.delete_note(nid)


@app.get("/api/projects/{pid}/versions")
def project_versions(pid: str):
    return {"versions": projects_store.list_versions(pid)}


@app.post("/api/projects/{pid}/versions")
def project_version_save(pid: str, req: VersionSaveIn):
    v = projects_store.save_version(pid, req.state, req.label, req.pinned, req.auto)
    return {"ok": True, "version": v}


@app.get("/api/projects/{pid}/versions/{vid}")
def project_version_get(pid: str, vid: str):
    v = projects_store.get_version(vid)
    if not v:
        return Response(status_code=404, content="not found")
    return v


@app.patch("/api/projects/{pid}/versions/{vid}")
def project_version_patch(pid: str, vid: str, req: VersionPatchIn):
    patch = {}
    if req.label is not None:
        patch["label"] = req.label
    if req.pinned is not None:
        patch["pinned"] = req.pinned
    return projects_store.update_version(vid, patch)


@app.delete("/api/projects/{pid}/versions/{vid}")
def project_version_delete(pid: str, vid: str):
    return projects_store.delete_version(vid)



if EBOOKS.exists():
    app.mount("/ebooks", StaticFiles(directory=str(EBOOKS), html=True), name="ebooks")
DECK_OUT = EBOOK_HTML / "_deck_out"
DECK_OUT.mkdir(parents=True, exist_ok=True)
app.mount("/deck_out", StaticFiles(directory=str(DECK_OUT)), name="deck_out")
DIST = EBOOK_HTML / "dist"
if DIST.exists():
    app.mount("/", StaticFiles(directory=str(DIST), html=True), name="app")
